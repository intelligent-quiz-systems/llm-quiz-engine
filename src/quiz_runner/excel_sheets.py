"""
Excel sheet builders for quiz test series analysis.
Each function adds one sheet to the given Workbook.
"""
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Colours ───────────────────────────────────────────────────────────────────
_HDR_BG   = "2F5496"
_HDR_FG   = "FFFFFF"
_OK_BG    = "E2EFDA"
_FAIL_BG  = "FCE4D6"
_WARN_BG  = "FFF2CC"
_NOTE_FG  = "808080"

NA = "n/a"


# ── Column dictionary ─────────────────────────────────────────────────────────
COLUMN_DICT = [
    ("test_id",                   "Identyfikator testu (nazwa folderu)"),
    ("label",                     "Opcjonalna etykieta testu"),
    ("topic",                     "Temat quizu"),
    ("difficulty",                "Poziom trudności: easy / medium / hard"),
    ("count_requested",           "Żądana liczba pytań"),
    ("count_generated",           "Wygenerowana liczba pytań"),
    ("ok",                        "True jeśli test zakończony sukcesem"),
    ("status",                    "Status: completed / failed / partial / dry-run"),
    ("duration_seconds",          "Czas trwania całego testu (s)"),
    ("avg_seconds_per_question",  "Średni czas na jedno wygenerowane pytanie (s)"),
    ("n_batches",                 "Liczba zaakceptowanych wsadów API"),
    ("n_rejected_attempts",       "Liczba odrzuconych prób generacji"),
    ("min_batch_size",            "Minimalny rozmiar wsadu osiągnięty podczas fallback"),
    ("total_tokens",              "Łączne tokeny (input + output) ze wszystkich wsadów"),
    ("total_input_tokens",        "Łączne tokeny wejściowe (prompt)"),
    ("total_output_tokens",       "Łączne tokeny wyjściowe (odpowiedź)"),
    ("total_reasoning_tokens",    "Łączne tokeny reasoning"),
    ("avg_tokens_per_question",   "Średnia tokenów (total) na wygenerowane pytanie"),
    ("n_quality_issues",          "Liczba problemów jakości wykrytych w wygenerowanym quizie"),
    ("total_rejection_duration_s","Łączny czas stracony na odrzucone próby (s)"),
    ("error",                     "Komunikat błędu jeśli status != completed"),
    ("batch_number",              "Numer wsadu w ramach testu (1-based)"),
    ("accepted_batch_size",       "Rozmiar zaakceptowanego wsadu"),
    ("accepted_questions",        "Pytania zaakceptowane w tym wsadzie"),
    ("total_attempts",            "Łączna liczba prób dla tego wsadu (zaakceptowane + odrzucone)"),
    ("attempt_duration_seconds",  "Czas trwania próby w sekundach"),
    ("input_tokens",              "Tokeny wejściowe dla wsadu"),
    ("output_tokens",             "Tokeny wyjściowe dla wsadu"),
    ("reasoning_tokens",          "Tokeny reasoning dla wsadu"),
    ("system_prompt_chars",       "Długość system prompt w znakach"),
    ("user_prompt_chars",         "Długość user prompt w znakach"),
    ("total_prompt_chars",        "Łączna długość promptu w znakach"),
    ("guardrail_question_count",  "Liczba pytań w kontekście guardrail"),
    ("guardrail_chars",           "Długość tekstu guardrail w znakach"),
    ("decision",                  "Decyzja po nieudanej próbie: reduce / retry / halt"),
    ("error_type",                "Typ wyjątku Python"),
    ("attempted_batch_size",      "Rozmiar wsadu w momencie odrzucenia"),
    ("provider_error_info",       "Szczegóły błędu od providera LLM (jako string)"),
    ("issue_type",                "Typ problemu jakości: duplicate / similar / duplicate_options"),
    ("question_index",            "Indeks pytania (0-based)"),
    ("question_text",             "Tekst pytania"),
    ("detail",                    "Szczegół problemu jakości"),
]


# ── Style helpers ─────────────────────────────────────────────────────────────

def _hdr(ws, row: int, col: int, value: str):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True, color=_HDR_FG, size=10)
    c.fill = PatternFill("solid", fgColor=_HDR_BG)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    return c


def _auto_width(ws, min_w: int = 10, max_w: int = 45) -> None:
    for col_cells in ws.columns:
        width = max(
            (len(str(c.value)) if c.value is not None else 0 for c in col_cells),
            default=min_w,
        )
        ltr = get_column_letter(col_cells[0].column)
        ws.column_dimensions[ltr].width = min(max(width + 2, min_w), max_w)


def _write_rows(ws, headers: list[str], rows: list[dict], start_row: int = 2) -> None:
    for ri, r in enumerate(rows, start=start_row):
        for col, h in enumerate(headers, 1):
            val = r.get(h, NA)
            if isinstance(val, dict):
                val = str(val)
            ws.cell(row=ri, column=col, value=val)


def _round(v, digits: int = 3):
    try:
        return round(float(v), digits) if v is not None else NA
    except (TypeError, ValueError):
        return NA


def _pct(numerator, denominator, digits: int = 1):
    try:
        return round(100.0 * numerator / denominator, digits) if denominator else NA
    except (TypeError, ValueError):
        return NA


# ── Sheet builders ────────────────────────────────────────────────────────────

"""
Event warehouse column specification: (en_name, pl_description).
One row per event: accepted_batch, rejected_attempt, quality_issue, generation_error.
"""
_PIVOT_COLUMNS: list[tuple[str, str]] = [
    # ── Run / test context ────────────────────────────────────────────────────
    ("series_group",           "Seria testów"),
    ("run_datetime",           "Data i czas serii"),
    ("git_branch",             "Gałąź git"),
    ("nr_global",              "Nr globalny testu"),
    ("label",                  "Etykieta testu"),
    ("test_id",                "ID testu"),
    ("topic",                  "Temat quizu"),
    ("difficulty",             "Trudność"),
    ("target_questions",       "Cel: pytania"),
    ("accepted_questions",     "Zaakceptowanych pytań"),
    ("missing_questions",      "Brakujące pytania"),
    ("completion_pct",         "Kompletność (%)"),
    ("status",                 "Status testu"),
    ("duration_seconds",       "Czas testu (s)"),
    ("n_batches",              "Batche łącznie"),
    ("n_rejected_attempts",    "Odrzucone próby (test)"),
    ("min_batch_size",         "Min. rozmiar wsadu"),
    ("has_rejection",          "Były odrzucenia"),
    # ── Event ─────────────────────────────────────────────────────────────────
    ("event_type",             "Typ zdarzenia EN"),
    ("event_type_pl",          "Typ zdarzenia PL"),
    ("event_seq",              "Nr zdarzenia w teście"),
    ("batch_number",           "Nr wsadu (batcha)"),
    # ── Accepted batch ────────────────────────────────────────────────────────
    ("batch_duration_seconds", "Czas wsadu (s)"),
    ("attempts_count",         "Próby w wsadzie"),
    ("rejected_attempts_count","Odrzucone w wsadzie"),
    ("accepted_questions_count","Pytania zaakceptowane w wsadzie"),
    ("final_accepted_batch_size","Ostateczny rozmiar wsadu"),
    ("accepted_attempt_batch_size","Rozmiar zaakceptowanej próby"),
    ("accepted_attempt_duration_seconds","Czas zaakceptowanej próby (s)"),
    ("gross_s_per_q",          "Brutto s/pyt"),
    ("clean_s_per_q",          "Netto s/pyt"),
    ("rejection_overhead_s_total","Narzut odrzuceń (s)"),
    ("rejection_overhead_s_per_q","Narzut odrzuceń/pyt (s)"),
    ("rejection_overhead_pct", "Narzut odrzuceń (%)"),
    # ── Rejected attempt ──────────────────────────────────────────────────────
    ("attempt_number",         "Nr próby"),
    ("attempted_batch_size",   "Rozmiar próby"),
    ("attempt_duration_seconds","Czas próby (s)"),
    ("reject_reason",          "Powód odrzucenia EN"),
    ("reject_reason_pl",       "Powód odrzucenia PL"),
    ("reject_family",          "Rodzina błędu"),
    ("json_failure_kind",      "Rodzaj błędu JSON"),
    ("json_failure_kind_pl",   "Rodzaj błędu JSON PL"),
    ("provider_error_message_type","Typ komunikatu błędu"),
    ("provider_http_status",   "HTTP status"),
    ("provider_error_code",    "Kod błędu providera"),
    ("provider_error_type",    "Typ błędu providera"),
    ("provider_error_message", "Komunikat błędu providera"),
    ("provider_payload_present","Payload obecny"),
    ("provider_failed_generation_present","Nieudana generacja obecna"),
    ("provider_failed_generation_blank","Generacja pusta"),
    # ── Tokens ────────────────────────────────────────────────────────────────
    ("token_usage_known",      "Tokeny znane"),
    ("output_tokens_wasted",   "Tokeny zmarnowane"),
    ("output_token_usage_pct", "Tokeny wyjściowe (%)"),
    ("input_tokens_estimate",  "Tokeny wejściowe (est.)"),
    ("input_tokens",           "Tokeny wejściowe"),
    ("output_tokens",          "Tokeny wyjściowe"),
    ("reasoning_tokens",       "Tokeny reasoning"),
    ("total_tokens",           "Tokeny łącznie"),
    # ── Guardrail / prompt ────────────────────────────────────────────────────
    ("guardrail_question_count","Pytania guardrail"),
    ("guardrail_chars",        "Znaki guardrail"),
    ("total_prompt_chars",     "Znaki promptu łącznie"),
    ("failed_generation_chars","Znaki nieudanej generacji"),
    # ── Test-level token summary ──────────────────────────────────────────────
    ("n_batches",              "Batche (test) — duplikat kontekstu"),
    ("n_rejected_attempts",    "Odrzucone (test) — duplikat kontekstu"),
    ("min_batch_size",         "Min wsad (test) — duplikat kontekstu"),
    ("total_tokens_known",     "Tokeny znane (test)"),
    ("tokens_per_target_q",    "Tokeny / cel. pytanie"),
    ("rejected_time_s_detailed","Czas odrzuceń szczeg. (s)"),
    # ── Normalized event fields ───────────────────────────────────────────────
    ("event_duration_seconds", "Czas zdarzenia (s)"),
    ("event_batch_size",       "Rozmiar zdarzenia (wsad)"),
]

# Remove the duplicated context columns that appear twice
_SEEN: set[str] = set()
_PIVOT_COLUMNS_DEDUP: list[tuple[str, str]] = []
for _en, _pl in _PIVOT_COLUMNS:
    if _en not in _SEEN:
        _PIVOT_COLUMNS_DEDUP.append((_en, _pl))
        _SEEN.add(_en)
_PIVOT_COLUMNS = _PIVOT_COLUMNS_DEDUP

# Event type colours
_EVENT_FILL = {
    "accepted_batch":   "E2EFDA",  # green
    "rejected_attempt": "FCE4D6",  # red/orange
    "quality_issue":    "FFF2CC",  # yellow
    "generation_error": "F4CCCC",  # dark red
}


def add_pivot_full_sheet(wb: Workbook, event_rows: list[dict]) -> None:
    """
    Event warehouse fact table.
    Two-row bilingual header: row 1 = Polish description (rotated 90),
    row 2 = English machine name (rotated 90). Data from row 3.
    """
    ws = wb.create_sheet("Dane_pivot_FULL")

    _HDR_BG2 = "1F3864"   # darker blue for EN row

    # Row 1: Polish description — rotated
    for col, (en_name, pl_name) in enumerate(_PIVOT_COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=pl_name)
        c.font      = Font(bold=True, color=_HDR_FG, size=9)
        c.fill      = PatternFill("solid", fgColor=_HDR_BG)
        c.alignment = Alignment(textRotation=90, wrap_text=False,
                                vertical="bottom", horizontal="center")

    # Row 2: English machine name — rotated, slightly darker
    for col, (en_name, pl_name) in enumerate(_PIVOT_COLUMNS, 1):
        c = ws.cell(row=2, column=col, value=en_name)
        c.font      = Font(bold=False, color=_HDR_FG, size=8, italic=True)
        c.fill      = PatternFill("solid", fgColor=_HDR_BG2)
        c.alignment = Alignment(textRotation=90, wrap_text=False,
                                vertical="bottom", horizontal="center")

    ws.row_dimensions[1].height = 110
    ws.row_dimensions[2].height = 90

    # Data rows from row 3
    cols = [en for en, _ in _PIVOT_COLUMNS]
    for ri, row in enumerate(event_rows, start=3):
        ev_type = row.get("event_type", "")
        fill_color = _EVENT_FILL.get(ev_type)
        for col, key in enumerate(cols, 1):
            val = row.get(key)
            if isinstance(val, dict):
                val = str(val)
            cell = ws.cell(row=ri, column=col, value=val)
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)

    # Freeze at row 3 (below both header rows)
    ws.freeze_panes = "A3"
    # Auto-filter on header row 2
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}2"

    # Column widths: narrow for numeric, wider for text
    _TEXT_COLS = {"series_group", "test_id", "topic", "label", "status",
                  "event_type", "event_type_pl", "reject_reason_pl",
                  "provider_error_message", "reject_family"}
    for col, (en_name, _) in enumerate(_PIVOT_COLUMNS, 1):
        ltr = get_column_letter(col)
        ws.column_dimensions[ltr].width = 18 if en_name in _TEXT_COLS else 8


def add_dashboard_sheet(
    wb: Workbook,
    results: list[dict],
    enriched_results: list[dict],
    batch_rows: list[dict],
    rejection_rows: list[dict],
) -> None:
    ws = wb.create_sheet("01_Dashboard")
    completed   = sum(1 for r in results if r.get("ok"))
    failed      = len(results) - completed
    total_dur   = sum(r.get("duration_seconds") or 0 for r in results)
    total_tok   = sum((r.get("total_tokens") or 0) for r in results) or None
    total_rej   = sum((r.get("n_rejected_attempts") or 0) for r in results)
    total_rej_dur = sum((r.get("total_rejection_duration_s") or 0) for r in enriched_results)
    has_diag    = any(r.get("n_batches") is not None for r in results)
    n_batches_all = sum((r.get("n_batches") or 0) for r in results)

    ws.cell(row=1, column=1, value="Quiz Test Series — Summary").font = Font(bold=True, size=12)
    rows_data = [
        ("Generated at",              datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Total tests",               len(results)),
        ("Completed",                 completed),
        ("Failed",                    failed),
        ("Success rate",              _pct(completed, len(results)) if results else NA),
        ("Total duration (s)",        _round(total_dur, 1)),
        ("Total tokens (all batches)", total_tok if total_tok else NA),
        ("Total API batches",         n_batches_all if has_diag else NA),
        ("Total rejected attempts",   total_rej if has_diag else NA),
        ("Rejection overhead (s)",    _round(total_rej_dur, 1) if has_diag else NA),
        ("Diagnostics available",     "yes" if has_diag else "no"),
    ]
    for i, (lbl, val) in enumerate(rows_data, start=3):
        ws.cell(row=i, column=1, value=lbl).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30


def add_testy_sheet(wb: Workbook, enriched_results: list[dict]) -> None:
    ws = wb.create_sheet("02_Testy")
    headers = [
        "test_id", "topic", "difficulty",
        "count_requested", "count_generated", "status",
        "duration_seconds", "avg_seconds_per_question",
        "n_batches", "n_rejected_attempts", "min_batch_size",
        "total_tokens", "avg_tokens_per_question",
        "n_quality_issues", "total_rejection_duration_s", "error",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)

    for ri, r in enumerate(enriched_results, start=2):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=ri, column=col, value=r.get(h, NA))
            if r.get("ok"):
                cell.fill = PatternFill("solid", fgColor=_OK_BG)
            elif r.get("status") == "failed":
                cell.fill = PatternFill("solid", fgColor=_FAIL_BG)

    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_podsumowania_sheet(wb: Workbook, results: list[dict]) -> None:
    """Aggregated summaries per topic and per difficulty."""
    ws = wb.create_sheet("03_Podsumowania")

    def _agg(rows: list[dict], group_key: str) -> list[dict]:
        groups: dict[str, list] = defaultdict(list)
        for r in rows:
            groups[r.get(group_key, NA)].append(r)
        out = []
        for key, items in sorted(groups.items()):
            n = len(items)
            ok = sum(1 for r in items if r.get("ok"))
            durs = [r.get("duration_seconds") for r in items if r.get("duration_seconds") is not None]
            toks = [r.get("total_tokens") for r in items if r.get("total_tokens") is not None]
            rejs = [r.get("n_rejected_attempts") or 0 for r in items]
            batches = [r.get("n_batches") or 0 for r in items]
            out.append({
                group_key:          key,
                "n_tests":          n,
                "n_completed":      ok,
                "success_rate_pct": _pct(ok, n),
                "avg_duration_s":   _round(sum(durs) / len(durs)) if durs else NA,
                "avg_tokens":       _round(sum(toks) / len(toks)) if toks else NA,
                "avg_n_rejected":   _round(sum(rejs) / n),
                "avg_n_batches":    _round(sum(batches) / n),
            })
        return out

    # Section: per topic
    ws.cell(row=1, column=1, value="Per topic").font = Font(bold=True, size=11)
    topic_headers = ["topic", "n_tests", "n_completed", "success_rate_pct",
                     "avg_duration_s", "avg_tokens", "avg_n_rejected", "avg_n_batches"]
    for col, h in enumerate(topic_headers, 1):
        _hdr(ws, 2, col, h)
    topic_rows = _agg(results, "topic")
    _write_rows(ws, topic_headers, topic_rows, start_row=3)

    # Section: per difficulty (below topic table with gap)
    gap_row = 3 + len(topic_rows) + 2
    ws.cell(row=gap_row, column=1, value="Per difficulty").font = Font(bold=True, size=11)
    diff_headers = ["difficulty", "n_tests", "n_completed", "success_rate_pct",
                    "avg_duration_s", "avg_tokens", "avg_n_rejected", "avg_n_batches"]
    for col, h in enumerate(diff_headers, 1):
        _hdr(ws, gap_row + 1, col, h)
    diff_rows = _agg(results, "difficulty")
    _write_rows(ws, diff_headers, diff_rows, start_row=gap_row + 2)

    _auto_width(ws)


def add_bledy_per_test_sheet(wb: Workbook, rejection_rows: list[dict], results: list[dict]) -> None:
    """Per-test rejection summary — only tests with at least one rejection."""
    ws = wb.create_sheet("04_Bledy_per_test")
    headers = [
        "test_id", "topic", "difficulty",
        "n_rejections", "dominant_decision", "dominant_error_type",
        "total_rejection_duration_s", "mean_rejection_duration_s",
        "n_halt", "n_reduce", "n_retry",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)

    # Build lookup: test_id → topic/difficulty
    meta = {r.get("test_id", ""): r for r in results}

    groups: dict[str, list] = defaultdict(list)
    for r in rejection_rows:
        groups[r.get("test_id", NA)].append(r)

    rows_out = []
    for test_id, items in sorted(groups.items()):
        decisions    = Counter(r.get("decision", NA) for r in items)
        error_types  = Counter(r.get("error_type", NA) for r in items)
        durs = [r.get("attempt_duration_seconds") for r in items if r.get("attempt_duration_seconds") is not None]
        result_meta  = meta.get(test_id, {})
        rows_out.append({
            "test_id":                    test_id,
            "topic":                      result_meta.get("topic", NA),
            "difficulty":                 result_meta.get("difficulty", NA),
            "n_rejections":               len(items),
            "dominant_decision":          decisions.most_common(1)[0][0] if decisions else NA,
            "dominant_error_type":        error_types.most_common(1)[0][0] if error_types else NA,
            "total_rejection_duration_s": _round(sum(durs)),
            "mean_rejection_duration_s":  _round(sum(durs) / len(durs)) if durs else NA,
            "n_halt":                     decisions.get("halt", 0),
            "n_reduce":                   decisions.get("reduce", 0),
            "n_retry":                    decisions.get("retry", 0),
        })

    for ri, r in enumerate(rows_out, start=2):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=ri, column=col, value=r.get(h, NA))
            if r.get("n_halt", 0) > 0:
                cell.fill = PatternFill("solid", fgColor=_FAIL_BG)

    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_kategorie_bledow_sheet(wb: Workbook, rejection_rows: list[dict]) -> None:
    """Aggregated rejection counts by decision × error_type."""
    ws = wb.create_sheet("05_Kategorie_bledow")
    headers = [
        "decision", "error_type", "n_attempts",
        "mean_duration_s", "mean_batch_size", "pct_of_all_rejections",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)

    groups: dict[tuple, list] = defaultdict(list)
    for r in rejection_rows:
        key = (r.get("decision", NA), r.get("error_type", NA))
        groups[key].append(r)

    total_rej = len(rejection_rows)
    rows_out = []
    for (decision, error_type), items in sorted(groups.items()):
        durs  = [x.get("attempt_duration_seconds") for x in items if x.get("attempt_duration_seconds") is not None]
        sizes = [x.get("attempted_batch_size") for x in items if x.get("attempted_batch_size") is not None]
        rows_out.append({
            "decision":              decision,
            "error_type":            error_type,
            "n_attempts":            len(items),
            "mean_duration_s":       _round(sum(durs) / len(durs)) if durs else NA,
            "mean_batch_size":       _round(sum(sizes) / len(sizes), 1) if sizes else NA,
            "pct_of_all_rejections": _pct(len(items), total_rej) if total_rej else NA,
        })

    _write_rows(ws, headers, rows_out)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_batche_srednie_sheet(wb: Workbook, batch_rows: list[dict], results: list[dict]) -> None:
    """Average batch statistics per test."""
    ws = wb.create_sheet("06_Batche_srednie")
    headers = [
        "test_id", "topic", "difficulty",
        "n_batches", "avg_batch_size", "avg_duration_s",
        "avg_input_tokens", "avg_output_tokens",
        "avg_reasoning_tokens", "avg_total_tokens",
        "avg_prompt_chars", "avg_guardrail_chars",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)

    meta = {r.get("test_id", ""): r for r in results}

    groups: dict[str, list] = defaultdict(list)
    for b in batch_rows:
        groups[b.get("test_id", NA)].append(b)

    rows_out = []
    for test_id, items in sorted(groups.items()):
        def _avg(key):
            vals = [x.get(key) for x in items if x.get(key) is not None]
            return _round(sum(vals) / len(vals)) if vals else NA

        result_meta = meta.get(test_id, {})
        rows_out.append({
            "test_id":              test_id,
            "topic":                result_meta.get("topic", NA),
            "difficulty":           result_meta.get("difficulty", NA),
            "n_batches":            len(items),
            "avg_batch_size":       _avg("accepted_batch_size"),
            "avg_duration_s":       _avg("attempt_duration_seconds"),
            "avg_input_tokens":     _avg("input_tokens"),
            "avg_output_tokens":    _avg("output_tokens"),
            "avg_reasoning_tokens": _avg("reasoning_tokens"),
            "avg_total_tokens":     _avg("total_tokens"),
            "avg_prompt_chars":     _avg("total_prompt_chars"),
            "avg_guardrail_chars":  _avg("guardrail_chars"),
        })

    _write_rows(ws, headers, rows_out)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_guardrail_sheet(wb: Workbook, batch_rows: list[dict]) -> None:
    """Per-batch guardrail context size (only rows where guardrail was active)."""
    ws = wb.create_sheet("07_Guardrail")
    headers = [
        "test_id", "batch_number",
        "guardrail_question_count", "guardrail_chars",
        "system_prompt_chars", "user_prompt_chars", "total_prompt_chars",
        "input_tokens",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)
    active = [r for r in batch_rows if r.get("guardrail_question_count")]
    _write_rows(ws, headers, active)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_raw_rejections_sheet(wb: Workbook, rejection_rows: list[dict]) -> None:
    ws = wb.create_sheet("08_Raw_rejections")
    headers = [
        "test_id", "batch_number", "attempt_number", "attempted_batch_size",
        "decision", "error_type", "attempt_duration_seconds",
        "system_prompt_chars", "user_prompt_chars", "total_prompt_chars",
        "guardrail_question_count", "guardrail_chars", "provider_error_info",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)
    for ri, r in enumerate(rejection_rows, start=2):
        for col, h in enumerate(headers, 1):
            val = r.get(h, NA)
            if isinstance(val, dict):
                val = str(val)
            cell = ws.cell(row=ri, column=col, value=val)
            if r.get("decision") == "halt":
                cell.fill = PatternFill("solid", fgColor=_FAIL_BG)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_raw_batches_sheet(wb: Workbook, batch_rows: list[dict]) -> None:
    ws = wb.create_sheet("09_Raw_batches")
    headers = [
        "test_id", "batch_number", "accepted_batch_size", "accepted_questions",
        "total_attempts", "attempt_duration_seconds",
        "input_tokens", "output_tokens", "reasoning_tokens", "total_tokens",
        "system_prompt_chars", "user_prompt_chars", "total_prompt_chars",
        "guardrail_question_count", "guardrail_chars",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)
    _write_rows(ws, headers, batch_rows)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_quality_issues_sheet(wb: Workbook, quality_rows: list[dict]) -> None:
    ws = wb.create_sheet("10_Quality_issues")
    headers = [
        "test_id", "issue_type", "question_index", "question_text", "detail",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws, 1, col, h)
    _write_rows(ws, headers, quality_rows)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def add_slownik_kolumn_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("11_Slownik_kolumn")
    _hdr(ws, 1, 1, "Kolumna")
    _hdr(ws, 1, 2, "Opis")
    for ri, (col, desc) in enumerate(COLUMN_DICT, start=2):
        ws.cell(row=ri, column=1, value=col)
        ws.cell(row=ri, column=2, value=desc)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 62
